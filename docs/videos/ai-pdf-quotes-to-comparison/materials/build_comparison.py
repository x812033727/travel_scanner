#!/usr/bin/env python3
"""Verify a real AI extraction, then calculate and export comparison files.

The AI extraction JSON is an input, never invented by this program. Fixture
truth is used only in the independent verification pass before output.
"""
from __future__ import annotations
import argparse
import csv
import hashlib
import json
import re
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MATERIALS = Path(__file__).resolve().parent
MONEY_FIELDS = {'unit_price', 'goods_tax_rate', 'shipping_amount', 'shipping_tax_rate'}
REQUIRED_FIELDS = {'specification','quoted_quantity','quantity_min','quantity_max','currency','unit_price','unit_tax_mode','goods_tax_rate','shipping_amount','shipping_tax_mode','shipping_tax_rate','lead_business_days'}
CENT = Decimal('0.01')

def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()

def folded(text: str) -> str:
    return ' '.join(text.split()).casefold()

def money(value: Decimal) -> str:
    return str(value.quantize(CENT, rounding=ROUND_HALF_UP))

def compare_value(field: str, actual, expected) -> bool:
    if actual is None or expected is None:
        return actual is expected
    if field in MONEY_FIELDS:
        return Decimal(str(actual)) == Decimal(str(expected))
    return actual == expected

def delivered(fields: dict, quantity: int) -> dict:
    value = lambda key: fields[key]['value']
    if not isinstance(quantity, int) or isinstance(quantity, bool) or quantity < 1:
        raise ValueError('Quantity must be a positive integer')
    if not value('quantity_min') <= quantity <= value('quantity_max'):
        raise ValueError('Quantity outside explicitly quoted range')
    goods = Decimal(value('unit_price')) * quantity
    if value('unit_tax_mode') == 'exclusive':
        if value('goods_tax_rate') is None:
            raise ValueError('Unknown goods tax rate')
        goods *= Decimal(1) + Decimal(value('goods_tax_rate'))
    elif value('unit_tax_mode') != 'inclusive':
        raise ValueError('Unknown goods tax basis')
    shipping = value('shipping_amount')
    if shipping is None or value('shipping_tax_mode') == 'unknown':
        return {'goods_gross': money(goods), 'shipping_gross': None, 'delivered_total': None, 'status': '待確認：運費及其稅額未報', 'comparable': False}
    shipping_gross = Decimal(shipping)
    if value('shipping_tax_mode') == 'exclusive':
        if value('shipping_tax_rate') is None:
            raise ValueError('Unknown shipping tax rate')
        shipping_gross *= Decimal(1) + Decimal(value('shipping_tax_rate'))
    elif value('shipping_tax_mode') != 'inclusive':
        raise ValueError('Unknown shipping tax basis')
    return {'goods_gross': money(goods), 'shipping_gross': money(shipping_gross), 'delivered_total': money(goods + shipping_gross), 'status': '可比較', 'comparable': True}

def make_rows(quote_map: dict, files: list[str], qty: int) -> list[dict]:
    rows = []
    for filename in files:
        quote = quote_map[filename]; fields = quote['fields']
        result = delivered(fields, qty)
        rows.append({'supplier': quote['supplier'], 'source_file': filename, 'quantity': qty,
          'unit_price': fields['unit_price']['value'], 'unit_tax_mode': fields['unit_tax_mode']['value'],
          **result, 'lead_business_days': fields['lead_business_days']['value'],
          'price_source': f"{filename} p.{fields['unit_price']['page']}",
          'shipping_source': f"{filename} p.{fields['shipping_amount']['page']}"})
    numeric = [Decimal(row['delivered_total']) for row in rows if row['comparable']]
    best = min(numeric) if numeric else None
    for row in rows:
        row['lowest_comparable'] = row['comparable'] and Decimal(row['delivered_total']) == best
    return rows

def write_csv(target: Path, rows: list[dict]) -> None:
    with target.open('w', newline='', encoding='utf-8-sig') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader(); writer.writerows(rows)

def write_workbook(target: Path, scenarios: dict, quotes: list[dict]) -> None:
    wb = Workbook(); readme = wb.active; readme.title = '使用說明'
    for line in [
      ['原創虛構 PDF 報價練習', '所有供應商與條件均屬虛構'],
      ['計算', 'Python Decimal 獨立計算；本檔寫入已算值，不宣稱 openpyxl 會計算公式'],
      ['稅率', '5% 只是本練習明示的計算假設，不是實際稅法或稅務建議'],
      ['缺漏', '未報運費保留空白及待確認，不填零，也不參與總價排序'],
      ['比較', '總價最低只在符合規格且資料完整的候選中成立'],
      ['修訂', 'cedar-revised.pdf 取代 cedar.pdf；不是第 4 家供應商'],
      ['數量', '僅 10–20 張是三家書面報價適用範圍，超出需重詢'],
      ['來源', '逐欄來源工作表保留檔名、頁碼及本次 AI 抽取的原文證據'],
      ['可編輯', '儲存格可編輯；若變更輸入，請重新執行程式產檔；不是有即時公式的計算器']
    ]: readme.append(line)
    headers = ['供應商','原檔','數量','原單價','單價稅基','商品含明示稅額','運費含明示稅額','送達總價','狀態','交期工作天','單價來源','運費來源','已知可比最低']
    keys = ['supplier','source_file','quantity','unit_price','unit_tax_mode','goods_gross','shipping_gross','delivered_total','status','lead_business_days','price_source','shipping_source','lowest_comparable']
    labels = {'original-10':'原報價10張','revised-10':'修訂報價10張','revised-20':'修訂報價20張'}
    for scenario, rows in scenarios.items():
        ws = wb.create_sheet(labels.get(scenario, scenario)); ws.append(headers)
        for row in rows:
            vals = [row[key] for key in keys]
            for i, key in enumerate(keys):
                if key in {'unit_price','goods_gross','shipping_gross','delivered_total'} and vals[i] is not None:
                    vals[i] = float(Decimal(vals[i]))
            ws.append(vals)
            if row['lowest_comparable']:
                for cell in ws[ws.max_row]: cell.fill = PatternFill('solid', fgColor='DDF3E4')
            if not row['comparable']:
                for cell in ws[ws.max_row]: cell.fill = PatternFill('solid', fgColor='FFF0CC')
        ws.auto_filter.ref = ws.dimensions; ws.freeze_panes = 'C2'
    evidence = wb.create_sheet('逐欄來源'); evidence.append(['檔名','供應商','修訂','欄位','抽出值','原頁','原文證據'])
    for quote in quotes:
        for name, field in quote['fields'].items():
            evidence.append([quote['source_file'],quote['supplier'],quote['revision'],name,field['value'],field['page'],field['evidence']])
    for ws in wb:
        ws.row_dimensions[1].height = 29
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='123D52'); cell.font = Font(color='FFFFFF', bold=True)
        for row in ws:
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if isinstance(cell.value, float): cell.number_format = '#,##0.00'
        for col in range(1,ws.max_column+1):
            ws.column_dimensions[get_column_letter(col)].width = 23 if col not in {2,7} else 43
        if ws.title == '逐欄來源': ws.column_dimensions['G'].width = 90; ws.column_dimensions['E'].width = 65
        if ws.title == '使用說明': ws.column_dimensions['A'].width = 25; ws.column_dimensions['B'].width = 105
        ws.sheet_view.showGridLines = False
    wb.save(target)
    reread = load_workbook(target, data_only=True)
    assert reread['原報價10張']['H4'].value is None
    assert reread['修訂報價20張']['H4'].value == 3937.5
    assert not any(isinstance(c.value, str) and c.value.startswith('=') for ws in reread for row in ws for c in row)

def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument('--analysis', type=Path, required=True)
    parser.add_argument('--extracted', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    parser.add_argument('--request', type=Path, help='Read actual quantity and original/revised choice from a request JSON')
    args = parser.parse_args()
    analysis = json.loads(args.analysis.read_text()); extracted = json.loads(args.extracted.read_text())
    truth = json.loads((MATERIALS / 'fixture-truth.json').read_text())
    pages = {d['source_file']: {p['page']: p['text'] for p in d['pages']} for d in extracted['documents']}
    quotes = analysis['quotes']; quote_map = {q['source_file']: q for q in quotes}
    assert len(quotes) == 4 and set(quote_map) == set(truth['expected_rows']), 'Missing/duplicate/unexpected quotation'
    checks = []
    for filename, expected in truth['expected_rows'].items():
        quote = quote_map[filename]; fields = quote['fields']
        assert quote['supplier'] == expected['supplier']
        assert quote['revision'] == ('revised' if filename == 'cedar-revised.pdf' else 'original')
        assert set(fields) == REQUIRED_FIELDS
        for name, field in fields.items():
            page = field['page']; evidence = field['evidence']
            assert page in pages[filename], f'Invalid page {filename}:{name}'
            assert isinstance(evidence,str) and evidence.strip(), f'Missing evidence {filename}:{name}'
            assert folded(evidence) in folded(pages[filename][page]), f'Evidence not in actual page {filename}:{name}'
            if name in expected:
                assert compare_value(name, field['value'], expected[name]), f'Independent fixture mismatch {filename}:{name}'
        assert fields['currency']['value'] == 'TWD'
        assert fields['quoted_quantity']['value'] == 10
        spec = folded(fields['specification']['value'])
        for token in ['a6','105','148','350','matte','white','double','colour','identical']:
            assert token in spec, f'Incomplete specification {filename}:{token}'
        for qty in [10,20]:
            result = delivered(fields, qty)
            assert result['delivered_total'] == truth['expected_delivered_totals'][str(qty)][filename]
            checks.append({'check':'independent_total','file':filename,'quantity':qty,'actual':result['delivered_total'],'pass':True})
        try:
            delivered(fields,21)
        except ValueError:
            checks.append({'check':'reject_out_of_range_quantity_21','file':filename,'pass':True})
        else:
            raise AssertionError('Out-of-range quantity accepted')
    originals = ['aster.pdf','birch.pdf','cedar.pdf']; revised = ['aster.pdf','birch.pdf','cedar-revised.pdf']
    scenarios = {'original-10':make_rows(quote_map,originals,10),'revised-10':make_rows(quote_map,revised,10),'revised-20':make_rows(quote_map,revised,20)}
    for key, rows in scenarios.items():
        winners = [row['supplier'] for row in rows if row['lowest_comparable']]
        assert winners == truth['expected_winners'][key]
        assert len({row['supplier'] for row in rows}) == 3
        checks.append({'check':'winner_and_no_duplicate_revision','scenario':key,'winners':winners,'pass':True})
    assert delivered(quote_map['cedar.pdf']['fields'],10)['goods_gross'] == truth['known_goods_gross_original_cedar_10']
    assert not quote_map['aster.pdf']['questions'] and not quote_map['birch.pdf']['questions'] and not quote_map['cedar-revised.pdf']['questions'], 'Spurious missing terms for complete quotes'
    assert {q['field'] for q in quote_map['cedar.pdf']['questions']} >= {'shipping_amount','shipping_tax_rate'}, 'Missing genuine clarification questions'
    selected_request = None
    selected_rows = None
    if args.request:
        selected_request = json.loads(args.request.read_text())
        if set(selected_request) != {'scenario','quantity'} or selected_request['scenario'] not in {'original','revised'}:
            raise ValueError('Request must contain exactly scenario=original|revised and quantity')
        # Validate the requested quantity before any output is created.
        selected_rows = make_rows(quote_map, originals if selected_request['scenario']=='original' else revised, selected_request['quantity'])
    args.out.mkdir(parents=True,exist_ok=True)
    for key,rows in scenarios.items(): write_csv(args.out/f'{key}.csv',rows)
    baseline = sorted([{'supplier':quote_map[name]['supplier'], 'source_file':name, 'unit_price':quote_map[name]['fields']['unit_price']['value'], 'unit_price_source_page':quote_map[name]['fields']['unit_price']['page'], 'method':'Arithmetic sort of original headline unit prices only; not an AI failure demonstration', 'comparable_total_claimed':False} for name in originals],key=lambda row:Decimal(row['unit_price']))
    assert [row['supplier'] for row in baseline] == ['Cedar','Aster','Birch']
    write_csv(args.out/'headline-unit-sort-baseline.csv',baseline)
    baseline_note={'method':'Actual arithmetic unit-price sort of the three original quotes, without tax/delivery normalization','not_a_claimed_model_failure':True,'unit_price_order':[row['supplier'] for row in baseline],'unit_prices':[row['unit_price'] for row in baseline],'normalized_original_10_order':['Birch','Aster'],'original_Cedar':'not comparable: delivery cost and delivery tax are unknown','revised_10_groups':[['Birch'],['Aster','Cedar']],'revised_20_order':['Cedar','Aster','Birch']}
    (args.out/'baseline-and-normalized-order.json').write_text(json.dumps(baseline_note,ensure_ascii=False,indent=2)+'\n')
    checks.append({'check':'actual_headline_sort_is_Cedar_Aster_Birch_not_model_failure','pass':True})
    workbook_scenarios = dict(scenarios)
    if selected_request is not None:
        write_csv(args.out/'selected-comparison.csv',selected_rows)
        (args.out/'selected-comparison.json').write_text(json.dumps({'request':selected_request,'request_sha256':sha(args.request),'rows':selected_rows},ensure_ascii=False,indent=2)+'\n')
        workbook_scenarios[f"本次請求{selected_request['quantity']}張"] = selected_rows
        checks.append({'check':'actual_request_file_drives_selected_table','request_file':str(args.request),'request_sha256':sha(args.request),'quantity':selected_request['quantity'],'scenario':selected_request['scenario'],'totals':{row['supplier']:row['delivered_total'] for row in selected_rows},'pass':True})
    write_workbook(args.out/'comparison.xlsx',workbook_scenarios,quotes)
    (args.out/'comparison.json').write_text(json.dumps(scenarios,ensure_ascii=False,indent=2)+'\n')
    question_lines=['Cedar 原版 / cedar.pdf p.2']
    question_lines += [f"{i}. {question['question']}（欄位 {question['field']}；cedar.pdf p.{question['page']}）" for i,question in enumerate(quote_map['cedar.pdf']['questions'],1)]
    question_lines += ['', '修訂版已提供這些資料，因此 revised-10 與 revised-20 沒有此缺漏。']
    (args.out/'questions.txt').write_text('\n'.join(question_lines)+'\n')
    before=quote_map['cedar.pdf']['fields']; after=quote_map['cedar-revised.pdf']['fields']
    changes=[{'field':k,'before':before[k]['value'],'after':after[k]['value']} for k in REQUIRED_FIELDS if before[k]['value'] != after[k]['value']]
    assert {c['field'] for c in changes} == {'shipping_amount','shipping_tax_mode','shipping_tax_rate'}, 'Unexpected revision change'
    (args.out/'change-log.json').write_text(json.dumps({'replacement':'cedar-revised.pdf replaces cedar.pdf','changed_fields':sorted(changes,key=lambda c:c['field']),'preserved':'All goods/specification/unit price/tax rate/quantity-range/lead-time fields','original_total':None,'revised_10_total':'2100.00','quantity_20_recomputed_total':'3937.50'},ensure_ascii=False,indent=2)+'\n')
    lines=['# 原創 PDF 報價比較結果','','全部金額為虛構練習 TWD；5% 僅為本練習的明示計算假設。','']
    for key,rows in scenarios.items():
        lines.extend([f'## {key}','','|供應商|數量|商品含稅|運費含稅|送達總價|判斷|','|---|---:|---:|---:|---:|---|'])
        for row in rows:
            total=row['delivered_total'] if row['delivered_total'] is not None else '待確認'
            lines.append(f"|{row['supplier']}|{row['quantity']}|{row['goods_gross']}|{row['shipping_gross'] or '待確認'}|{total}|{row['status']}{'；已知可比最低' if row['lowest_comparable'] else ''}|")
        lines.append('')
    lines.extend(['原版 Cedar 只知道商品含稅 1837.50，並不知道送達總價。','修訂版與 Aster 在 10 張時同為 2100.00；Birch 2050.00 仍最低。','改成 20 張後：Aster 3990.00、Birch 4100.00、Cedar 修訂版 3937.50。','數量超出書面有效範圍 10–20 張時，程式拒絕計算，需重詢報價。','試算表寫入 Decimal 已算值；不假稱 openpyxl 會計算公式。'])
    (args.out/'comparison.md').write_text('\n'.join(lines)+'\n')
    report={'checked_at_utc':datetime.now(timezone.utc).isoformat(),'actual_ai_analyst':analysis['analyst'],'analysis_sha256':sha(args.analysis),'selected_request':selected_request,'request_sha256':sha(args.request) if args.request else None,'extracted_pages_sha256':sha(args.extracted),'truth_sha256':sha(MATERIALS/'fixture-truth.json'),'verification':'Independent fixture field assertions, exact page-evidence substring checks, Decimal totals, actual xlsx data_only readback, unknown shipping kept blank, revision replacement and out-of-range rejection','checks':checks,'all_passed':all(c['pass'] for c in checks),'formula_evaluation_claimed':False,'output_files':[{'name':p.name,'bytes':p.stat().st_size,'sha256':sha(p)} for p in sorted(args.out.iterdir()) if p.is_file() and p.name!='verification.json']}
    (args.out/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'output':str(args.out),'checks':len(checks),'all_passed':report['all_passed'],'winners':truth['expected_winners']},ensure_ascii=False))

if __name__ == '__main__': main()
