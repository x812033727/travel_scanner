#!/usr/bin/env python3
"""Public PDF -> reviewable outline -> citation-bound editable deck and PDF.

Original local workflow demonstration, not execution of Gemini Notebook or Canva.
"""
from datetime import datetime, timezone
import hashlib
import html
import json
from pathlib import Path
import subprocess
import time
import fitz
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

HERE = Path(__file__).resolve().parent
OUT = HERE / 'output'
PDF_URL = 'https://hai.stanford.edu/assets/files/ai_index_report_2026_chapter_4_economy.pdf'
INK, TEAL, RED, BG, MUTED = '#172d46', '#007f85', '#b83e2a', '#f8faf8', '#4b5a65'
FONT = 'Noto Sans CJK TC'

def write(name, obj):
    (HERE/name).write_text(json.dumps(obj, ensure_ascii=False, indent=2)+'\n')

def main():
    started = datetime.now(timezone.utc).isoformat(); clock=time.perf_counter()
    OUT.mkdir(exist_ok=True)
    src=HERE/'source-economy.pdf'; raw=src.read_bytes(); pdf_sha=hashlib.sha256(raw).hexdigest()
    doc=fitz.open(src); pages={p+1:doc[p].get_text() for p in [2,7,11,22]}
    write('source-selected-pages.json',pages)
    for p in [3,8,12,23]:
        doc[p-1].get_pixmap(matrix=fitz.Matrix(1,1)).save(str(OUT/f'source-pdf-page-{p:02}.png'))
    checks={'pages':len(doc),'global_corporate': '581.69' in pages[8], 'global_private':'344.66' in pages[8],
            'us_private':'285.88' in pages[12], 'china_private':'12.41' in pages[12], 'ai88':'88%' in pages[23],
            'summary70':'70%' in pages[3], 'detailed79':'79%' in pages[23]}
    assert all(v for k,v in checks.items() if k!='pages')
    records=[
      dict(id='C1',value=581.69,unit='billion USD',data_year=2025,metric='global corporate AI investment, four financing categories',pdf_page=8,printed_page=178,figure='4.2.1',status='supported'),
      dict(id='C2',value=344.66,unit='billion USD',data_year=2025,metric='global private AI investment, a subset of C1',pdf_page=8,printed_page=178,figure='4.2.1',status='supported'),
      dict(id='C3',value=285.88,unit='billion USD',data_year=2025,metric='US private AI investment; not all public/private AI spending',pdf_page=12,printed_page=182,figure='4.2.8',status='supported'),
      dict(id='C4',value=12.41,unit='billion USD',data_year=2025,metric='China private AI investment; not all public/private AI spending',pdf_page=12,printed_page=182,figure='4.2.8',status='supported'),
      dict(id='C5',value=88,unit='percent of survey respondents reporting organizational AI use',data_year=2025,metric='organization uses AI in at least one business function; self-reported survey, not population census',pdf_page=23,printed_page=193,figure='4.3.1',status='supported'),
      dict(id='C6',candidate_value=70,detailed_value=79,unit='percent',data_year=2025,status='unresolved',export_policy='annotation_only_not_adopted_rate',summary=dict(pdf_page=3,printed_page=173,wording='organizations; use generative AI in at least one business function'),detail=dict(pdf_page=23,printed_page=193,wording='respondents reporting their organizations regularly use generative AI in at least one business function'),explanation='Different wording/denominator description; exact question/cohort alignment not established. No claim that Stanford is wrong or that either rate is the current true adoption rate.')
    ]
    for r in records: r.update(source_url=PDF_URL, source_sha256=pdf_sha,report_year=2026)
    write('claim-ledger.json',records)
    brief={'creator':'Current conversation assistant, with deterministic local PDF extraction/rendering; not an independently called model or product UI', 'purpose':'3-minute team briefing (target, not measured speech duration)', 'sources':[PDF_URL], 'instruction':'先交可審大綱，逐項列來源頁、資料年、單位、母體及待確認處。未確認的數字不得畫成確定結論。回覆完全部修改後，才產生4頁簡報。', 'product_scope':'Gemini Notebook new Review feature is CITED official context only. No login or provider generation occurred.'}
    write('input-brief.json',brief)
    write('outline-before-review.json',{'kind':'Author-created planning draft based on the official summary, not claimed model output','slides':[{'title':'AI投資和採用概況','claims':['C1','C2','C3','C4','C5','C6']},{'title':'整理成一張總覽圖','claims':['C1','C2','C3','C4','C5','C6']}], 'candidate_genai_rate':70, 'not_rendered':True})
    write('review-feedback.json',{'all_feedback_before_render':True,'changes':['改成4頁，每頁一個讀者問題，先分投資額與使用率。','全球企業投資與其中私人投資不可相加；地域比較只比較私人融資，不稱全國AI總支出。','每張保留2025資料年、美元十億/受訪者母體，標2026報告與資料年的區別。','生成式AI採用率70/79的口徑對齊未解；只作兩處來源觀察，不當通用採用率結論。','加上每頁印刷頁碼、PDF頁序與圖號；產出可編輯PPTX及可看PDF。']})
    slide_info=[
      dict(id='S1',title='同一份報告，先問兩個問題',claims=[],notes='Editorial outline: money invested and self-reported organizational use are different metrics.'),
      dict(id='S2',title='2025：全球企業投資與其中私人投資',claims=['C1','C2'],pdf_pages=[8],printed_pages=[178],figures=['4.2.1']),
      dict(id='S3',title='2025：用同一私人融資口徑比較地區',claims=['C3','C4'],pdf_pages=[12],printed_pages=[182],figures=['4.2.8']),
      dict(id='S4',title='2025：受訪者回報組織使用AI',claims=['C5'],annotations=['C6'],pdf_pages=[3,23],printed_pages=[173,193],figures=['4.3.1'])]
    write('outline-after-review.json',slide_info)
    prs=Presentation(); prs.slide_width=Inches(16); prs.slide_height=Inches(9)
    image_files=[]; all_svg=[]

    for info in slide_info:
        slide=prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor.from_string(BG[1:])
        parts=[f'<svg xmlns="http://www.w3.org/2000/svg" width="1600" height="900" viewBox="0 0 1600 900"><title>{html.escape(info["title"])}</title><rect width="1600" height="900" fill="{BG}"/>']
        def rect(x,y,w,h,color):
            parts.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="{color}" rx="10"/>')
            sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x/100),Inches(y/100),Inches(w/100),Inches(h/100)); sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(color[1:]); sh.line.fill.background()
        def text(x,y,t,size=46,color=INK,bold=False):
            parts.append(f'<text x="{x}" y="{y}" font-family="{FONT},sans-serif" font-size="{size}" fill="{color}" font-weight="{700 if bold else 400}">{html.escape(t)}</text>')
            tb=slide.shapes.add_textbox(Inches(x/100),Inches((y-size)/100),Inches((1520-x)/100),Inches((size*1.5)/100)); tf=tb.text_frame; tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.word_wrap=False
            p=tf.paragraphs[0]; p.text=t; p.font.name=FONT; p.font.size=Pt(size*.72); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color[1:])
        text(80,80,'本地流程示範｜不是 Gemini Notebook 產品畫面',36,MUTED)
        text(80,165,info['title'],54,bold=True)
        rect(80,200,1440,3,'#bac8cf')
        if info['id']=='S1':
            rect(80,275,690,305,'#e2eeee'); rect(800,275,720,305,'#e5eaf3')
            text(120,360,'投入多少錢？',62,TEAL,True); text(120,445,'投資類型／地區／資料年',42)
            text(120,525,'單位：十億美元',46,bold=True)
            text(840,360,'多少人在回報使用？',58,INK,True); text(840,445,'受訪者／組織／業務功能',42)
            text(840,525,'單位：受訪者比例',46,bold=True)
            text(80,680,'先審問題與口徑，再把確認的數字放進簡報。',48,bold=True)
            text(80,795,'資料：2026 AI Index 經濟章；本例只展示其2025資料',36,MUTED)
        elif info['id']=='S2':
            text(80,280,'單位：十億美元｜企業總額包含多種融資活動',40,MUTED)
            for y,label,val,col in [(365,'全球企業投資',581.69,INK),(535,'其中：私人投資',344.66,TEAL)]:
                text(80,y,label,46,bold=True); rect(420,y-46,val/600*850,70,col); text(1320,y+5,f'{val:.2f}',46,bold=True)
            text(80,690,'私人投資是其中一部分，兩個數字不可相加。',46,bold=True)
            text(80,795,'來源：圖4.2.1｜報告p178／PDF第8頁｜Quid，2025',36,MUTED)
        elif info['id']=='S3':
            text(80,280,'單位：十億美元｜私人AI投資，不是全國AI總支出',40,MUTED)
            for y,label,val,col in [(365,'美國',285.88,INK),(535,'中國',12.41,TEAL)]:
                text(80,y,label,50,bold=True); rect(300,y-46,val/300*950,70,col); text(1300,y+5,f'{val:.2f}',46,bold=True)
            text(80,690,'同年份、同單位、同投資類型，才放在同一圖上。',46,bold=True)
            text(80,795,'來源：圖4.2.8｜報告p182／PDF第12頁｜Quid，2025',36,MUTED)
        else:
            text(80,280,'至少一項業務使用AI｜自述調查，不是人口普查',40,MUTED)
            rect(80,340,880,105,TEAL); text(1020,425,'88%',90,TEAL,True)
            text(80,530,'生成式AI數字暫不當作已確認的採用率：',44,bold=True)
            text(80,600,'摘要p173列70%；詳文p193列79%，措辭不同。',42,RED)
            text(80,665,'年份皆指2025；問法／分母是否一致仍待釐清。',42,RED)
            text(80,795,'來源：報告p173、p193／PDF第3、23頁｜圖4.3.1',36,MUTED)
        text(80,855,'本地原創排版與頁碼核對；沒有執行Notebook／Canva生成服務',32,MUTED)
        svg=OUT/(info['id'].lower()+'.svg'); svg.write_text('\n'.join(parts+['</svg>'])+'\n'); all_svg.append(svg)
        png=svg.with_suffix('.png'); subprocess.run(['ffmpeg','-hide_banner','-loglevel','error','-y','-i',str(svg),'-frames:v','1',str(png)],check=True); image_files.append(png)
    prs.save(OUT/'reviewed-briefing-editable.pptx')
    imgs=[Image.open(p).convert('RGB') for p in image_files]
    imgs[0].save(OUT/'reviewed-briefing.pdf',save_all=True,append_images=imgs[1:],resolution=100)
    # Check the substantive handoff, not just renderer success.
    assert {c for s in slide_info for c in s['claims']} == {r['id'] for r in records if r['status']=='supported'}
    assert 'C6' not in {c for s in slide_info for c in s['claims']}
    reopened=Presentation(OUT/'reviewed-briefing-editable.pptx'); assert len(reopened.slides)==4
    assert len(fitz.open(OUT/'reviewed-briefing.pdf'))==4
    citations={'source_pdf_url':PDF_URL,'source_sha256':pdf_sha,'page_number_convention':'pdf_page is1-based file position; printed_page is printed report page','slides':slide_info,'claims':records,'unresolved_policy':'C6 displayed only as two source observations, never adopted as an estimated rate','product_scope':brief['product_scope']}
    (OUT/'page-citation-manifest.json').write_text(json.dumps(citations,ensure_ascii=False,indent=2)+'\n')
    result={'status':'PASS_WITH_ONE_UNRESOLVED_SOURCE_ALIGNMENT','started_at_utc':started,'finished_at_utc':datetime.now(timezone.utc).isoformat(),'local_pipeline_seconds':round(time.perf_counter()-clock,3),'source_pdf_pages':len(doc),'source_numbers_found':checks,'supported_numeric_claims':5,'unresolved_rate_claims':1,'slides':4,'actual_outputs':['reviewed-briefing-editable.pptx','reviewed-briefing.pdf','s1.png','s2.png','s3.png','s4.png','page-citation-manifest.json'],'not_measured':['Notebook generation quality','Notebook account access','time saved vs another model','quota or API cost saved','true current GenAI adoption rate']}
    write('results.json',result)
    manifest={str(p.relative_to(HERE)):{'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(HERE.rglob('*')) if p.is_file() and p.name!='manifest.json'}
    write('manifest.json',{'created_at_utc':datetime.now(timezone.utc).isoformat(),'files':manifest})
    print(json.dumps(result,ensure_ascii=False,indent=2))

if __name__=='__main__': main()
